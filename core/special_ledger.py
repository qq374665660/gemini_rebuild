"""专项经费台账（外部立项 / 院自主立项）的解析逻辑。

两本台账的表头写“单位：万元”，实际填的是元；金额统一在此换算成万元。
"""

from __future__ import annotations

from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
import hashlib
import re
import unicodedata


LEDGER_EXTERNAL = 'external'
LEDGER_INSTITUTE = 'institute'
LEDGER_TYPE_CHOICES = [
    (LEDGER_EXTERNAL, '外部立项课题'),
    (LEDGER_INSTITUTE, '院自主立项课题'),
]
LEDGER_TYPE_LABELS = dict(LEDGER_TYPE_CHOICES)
SPECIAL_LEDGER_FORMAT_VERSION = 'special-ledger-v1'
SPECIAL_LEDGER_SHEET = '汇总表'
LEDGER_UNIT_DIVISOR = Decimal('10000')
# 台账单元格是 Excel 浮点公式结果（如 261491.019999999），元口径两位小数即为可信精度。
LEDGER_YUAN_QUANT = Decimal('0.01')

# 万元口径下的量级上限；超过即说明单位没换算或表结构变了。
LEDGER_TOTAL_CEILING_WAN = Decimal('1000000')
LEDGER_ROW_CEILING_WAN = Decimal('100000')

NAME_COLUMN = '课题名称'
STATUS_COLUMN = '研发进度'

# 每本台账的必填列，缺列直接拒绝而不是静默算出半张表。
LEDGER_REQUIRED_COLUMNS = {
    LEDGER_EXTERNAL: (NAME_COLUMN, STATUS_COLUMN, '已到账经费', '总执行额度', '可支出经费'),
    LEDGER_INSTITUTE: (NAME_COLUMN, STATUS_COLUMN, '预算额度', '总执行额度', '任务书剩余经费'),
}

# 台账列名 -> 模型字段。两本台账列位不同，只按表头文本定位。
LEDGER_TEXT_COLUMNS = {
    '归属单位': 'owning_unit',
    '经费来源': 'funder',
    '课题负责人': 'principal',
    '开始时间': 'start_text',
    '结束时间': 'end_text',
}
LEDGER_SEQUENCE_COLUMNS = {
    LEDGER_EXTERNAL: '经费明细',
    LEDGER_INSTITUTE: '台账链接',
}
LEDGER_MONEY_COLUMNS = {
    '课题合同经费': 'contract_total',
    '归属院/地下空间课题合同经费': 'contract_allocated',
    '已到账经费': 'received_amount',
    '预算额度': 'approved_budget',
    '总执行额度': 'executed_total',
    '可支出经费': 'remaining_amount',
    '任务书剩余经费': 'remaining_amount',
    '本年度可支配经费': 'year_disposable',
    '本年预算额': 'year_budget',
    '本年执行额度额度': 'year_executed',
    '本年执行额度': 'year_executed',
}
# 专项经费分母：一律取课题档案登记的专项经费，不取台账额度。
LEDGER_SPECIAL_BUDGET_FIELD = {
    LEDGER_EXTERNAL: 'external_funding',
    LEDGER_INSTITUTE: 'institute_funding',
}


class SpecialLedgerError(ValueError):
    """台账工作簿不符合专项经费监控格式。"""


def normalize_ledger_header(value):
    text = '' if value is None else unicodedata.normalize('NFKC', str(value))
    text = re.sub(r'[（(].*?[）)]', '', text)
    return re.sub(r'\s+', '', text).strip()


def normalize_project_name(value):
    """仅用于台账与系统课题的等值匹配，不剥离经费来源标记。"""
    text = '' if value is None else unicodedata.normalize('NFKC', str(value))
    text = re.sub(r'[《》“”‘’"\']', '', text)
    text = re.sub(r'\s+', '', text).lower()
    return re.sub(r'[^0-9a-z\u4e00-\u9fff]+', '', text)


def ledger_decimal(value):
    text = '' if value is None else unicodedata.normalize('NFKC', str(value))
    text = text.replace(',', '').replace('，', '').replace('¥', '').replace('￥', '').strip()
    if text in ('', '-', '—', '——', '/', 'nan', 'None', 'NaT'):
        return None
    negative = text.startswith('(') and text.endswith(')')
    if negative:
        text = text[1:-1]
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    return -amount if negative else amount


def file_digest(file_path):
    digest = hashlib.sha256()
    with open(file_path, 'rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def _worksheet(file_path):
    try:
        import openpyxl
    except Exception as exc:  # pragma: no cover - 仅服务部署缺依赖时触发
        raise SpecialLedgerError('服务端缺少 openpyxl，无法读取 Excel。') from exc

    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        raise SpecialLedgerError(f'无法打开台账文件：{exc}') from exc
    if SPECIAL_LEDGER_SHEET not in workbook.sheetnames:
        workbook.close()
        raise SpecialLedgerError(
            f'台账文件缺少“{SPECIAL_LEDGER_SHEET}”工作表，现有工作表：{"、".join(workbook.sheetnames) or "无"}'
        )
    return workbook, workbook[SPECIAL_LEDGER_SHEET]


def _locate_header(rows, ledger_type):
    """在前若干行里找到真正的表头行，跳过标题与多级表头。"""
    required = set(LEDGER_REQUIRED_COLUMNS[ledger_type])
    for offset, row in enumerate(rows[:12]):
        headers = {normalize_ledger_header(cell) for cell in row}
        if NAME_COLUMN in headers and required <= headers:
            return offset, row
    raise SpecialLedgerError(
        f'“{SPECIAL_LEDGER_SHEET}”里找不到{LEDGER_TYPE_LABELS[ledger_type]}的表头行，'
        f'必须包含列：{"、".join(LEDGER_REQUIRED_COLUMNS[ledger_type])}。'
    )


def parse_special_ledger(file_path, ledger_type, projects=(), assignment_lookup=None):
    """把一本台账的汇总表读成万元口径的行数据。

    assignment_lookup 形如 {(ledger_type, normalized_name): Project}，用于人工消解重名歧义。
    """
    if ledger_type not in LEDGER_TYPE_LABELS:
        raise SpecialLedgerError(f'未知的台账类型：{ledger_type}')
    assignment_lookup = assignment_lookup or {}

    workbook, worksheet = _worksheet(file_path)
    try:
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    header_offset, header_row = _locate_header(rows, ledger_type)
    column_index = {}
    for index, cell in enumerate(header_row):
        header = normalize_ledger_header(cell)
        if header and header not in column_index:
            column_index[header] = index

    name_index = column_index[NAME_COLUMN]
    status_index = column_index[STATUS_COLUMN]
    sequence_header = LEDGER_SEQUENCE_COLUMNS[ledger_type]
    sequence_index = column_index.get(sequence_header)

    records = []
    ignored_rows = []
    totals = {}

    for offset, row in enumerate(rows[header_offset + 1:], start=header_offset + 1):
        ledger_name = ('' if name_index >= len(row) or row[name_index] is None else str(row[name_index])).strip()
        if not ledger_name:
            continue
        status_text = ('' if status_index >= len(row) or row[status_index] is None else str(row[status_index])).strip()

        amounts = {}
        yuan_totals = {}
        for header, field_name in LEDGER_MONEY_COLUMNS.items():
            index = column_index.get(header)
            if index is None:
                continue
            value = row[index] if index < len(row) else None
            yuan = ledger_decimal(value)
            if yuan is None:
                continue
            yuan = yuan.quantize(LEDGER_YUAN_QUANT, rounding=ROUND_HALF_UP)
            yuan_totals[field_name] = yuan_totals.get(field_name, Decimal('0')) + yuan
        for field_name, yuan in yuan_totals.items():
            amounts[field_name] = yuan / LEDGER_UNIT_DIVISOR

        sequence_value = None
        if sequence_index is not None and sequence_index < len(row):
            sequence_value = ('' if row[sequence_index] is None else str(row[sequence_index])).strip()

        matched_project = None
        match_state = 'unmatched'
        normalized_name = normalize_project_name(ledger_name)
        candidates = [project for project in projects if normalize_project_name(project.name) == normalized_name] if normalized_name else []
        if normalized_name and assignment_lookup:
            assigned = assignment_lookup.get((ledger_type, normalized_name))
            if assigned is not None:
                matched_project = assigned
                match_state = 'manual'
        if matched_project is None:
            if len(candidates) == 1:
                matched_project = candidates[0]
                match_state = 'exact'
            elif len(candidates) > 1:
                # 系统里存在同名课题，钱挂错代价太大，交给人工确认。
                match_state = 'ambiguous'

        if matched_project is None and match_state != 'ambiguous':
            ignored_rows.append({
                'ledger_name': ledger_name,
                'ledger_status': status_text or '未填写',
                'funder': _text(row, column_index.get('经费来源')),
                'executed_total': amounts.get('executed_total'),
                'reason': '系统中没有同名课题',
            })
            continue

        records.append({
            'row_number': offset + 1,
            'sequence': (sequence_value or '').strip() or str(len(records) + 1),
            'project': matched_project,
            'match_state': match_state,
            'ledger_name': ledger_name,
            'ledger_status': status_text,
            'owning_unit': _text(row, column_index.get('归属单位')),
            'funder': _text(row, column_index.get('经费来源')),
            'principal': _text(row, column_index.get('课题负责人')),
            'start_text': _text(row, column_index.get('开始时间')),
            'end_text': _text(row, column_index.get('结束时间')),
            **amounts,
        })

    if not records:
        raise SpecialLedgerError(
            f'{LEDGER_TYPE_LABELS[ledger_type]}台账的“{SPECIAL_LEDGER_SHEET}”没有解析出任何课题行。'
        )

    # 汇总只统计已确认归属的行，避免重名歧义金额被随意计入某一课题后污染统计卡。
    confirmed = [record for record in records if record['match_state'] in {'exact', 'manual'}]
    for field_name in ('contract_total', 'contract_allocated', 'received_amount', 'approved_budget',
                       'executed_total', 'remaining_amount', 'year_disposable', 'year_budget', 'year_executed'):
        values = [record[field_name] for record in confirmed if record.get(field_name) is not None]
        totals[field_name] = sum(values, Decimal('0')) if values else None

    for field_name, total in totals.items():
        if total is not None and total > LEDGER_TOTAL_CEILING_WAN:
            raise SpecialLedgerError(
                f'{LEDGER_TYPE_LABELS[ledger_type]}台账的“{field_name}”合计 {total:,.2f} 万元，'
                '超出合理量级，疑似表头单位与数值不一致（台账表头写万元但填的是元），已拒绝导入。'
            )
    for record in records:
        for field_name, value in record.items():
            if isinstance(value, Decimal) and abs(value) > LEDGER_ROW_CEILING_WAN:
                raise SpecialLedgerError(
                    f'课题“{record["ledger_name"]}”的 {field_name} 为 {value:,.2f} 万元，超出合理量级，已拒绝导入。'
                )

    return {
        'ledger_type': ledger_type,
        'sheet_name': SPECIAL_LEDGER_SHEET,
        'header_row': header_offset + 1,
        'records': records,
        'ignored_rows': ignored_rows,
        'row_total': len(records) + len(ignored_rows),
        'matched_total': sum(1 for record in records if record['match_state'] in {'exact', 'manual'}),
        'ambiguous_total': sum(1 for record in records if record['match_state'] == 'ambiguous'),
        'ignored_total': len(ignored_rows),
        'totals': totals,
    }


def _text(row, index):
    if index is None or index >= len(row) or row[index] is None:
        return ''
    text = unicodedata.normalize('NFKC', str(row[index])).strip()
    return '' if text in ('nan', 'None', 'NaT') else text
