from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
import json
import tempfile
import zipfile
from unittest.mock import patch

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.messages.storage.fallback import FallbackStorage
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.template.loader import render_to_string
from django.test import RequestFactory, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .forms import ProjectForm
from .models import (
    APIConfig,
    MetricEvidence,
    MetricIndicatorDefinition,
    MetricsCategory,
    MetricsItem,
    OperationLog,
    Project,
    ProjectAnalysis,
    ExpenseImport,
    ExpenseSnapshot,
)
from .views import import_from_excel_view, extract_task_docx_view, query_assistant_view
from .docx_task_extractor import extract_task_pdf_text
from .analysis_parsing import parse_ai_analysis
from .ai_analysis import AIAnalysisService
from .query_assistant import answer_project_question
from .ai_query_assistant import (
    _tool_compare_project_years,
    _tool_find_projects,
    _tool_summarize_projects,
    answer_with_deepseek,
)
from .middleware import AccessControlMiddleware
from . import backup_schedule
from .expense_analysis import (
    EXPENSE_FORMAT_VERSION,
    SHORT_COMPANY_NAME,
    TARGET_COMPANIES,
    analyze_expense_workbook,
    canonicalize_expense_description,
    normalized_expense_description,
)


class ProjectFilterTests(TestCase):
    def setUp(self):
        self.admin_user = get_user_model().objects.create_user('filter-admin', password='StrongPass!234', is_staff=True)
        self.client.force_login(self.admin_user)
        Project.objects.create(
            project_id='P1',
            name='Alpha Project',
            ownership='GroupA',
            managing_unit='UnitA',
            level='LevelA',
            project_type='TypeA',
            role='Lead',
            start_year=2023,
            status='In Progress',
            project_lead='LeaderA',
            start_date=date(2023, 1, 1),
            planned_end_date=date(2024, 12, 31),
            total_budget=Decimal('120.00'),
            directory_path='P1',
        )
        Project.objects.create(
            project_id='P2',
            name='Beta Project',
            ownership='GroupB',
            managing_unit='UnitB',
            level='LevelB',
            project_type='TypeB',
            role='Member',
            start_year=2024,
            status='Completed',
            project_lead='LeaderB',
            start_date=date(2024, 1, 15),
            planned_end_date=date(2025, 12, 31),
            total_budget=Decimal('15.00'),
            directory_path='P2',
        )
        Project.objects.create(
            project_id='P3',
            name='Gamma Project',
            ownership='GroupA',
            managing_unit='UnitC',
            level='LevelA',
            project_type='TypeA',
            role='Lead',
            start_year=2024,
            status='In Progress',
            project_lead='LeaderC',
            start_date=date(2024, 5, 1),
            planned_end_date=date(2026, 1, 1),
            total_budget=Decimal('80.00'),
            directory_path='P3',
        )
        Project.objects.create(
            project_id='P4',
            name='Delta Project',
            ownership='GroupA',
            managing_unit='',
            level='LevelA',
            project_type='TypeA',
            role='Lead',
            start_year=2024,
            status='In Progress',
            project_lead='LeaderD',
            start_date=date(2024, 2, 1),
            planned_end_date=date(2026, 6, 1),
            total_budget=Decimal('60.00'),
            directory_path='P4',
        )

    def test_query_terms_and_filter(self):
        response = self.client.get(reverse('project_list'), {'q': 'Alpha UnitA'})
        self.assertEqual(response.status_code, 200)
        projects = list(response.context['projects'])
        self.assertEqual(len(projects), 1)
        self.assertEqual(projects[0].project_id, 'P1')

    def test_multi_select_and_budget_range(self):
        response = self.client.get(
            reverse('project_list'),
            {
                'year': ['2023', '2024'],
                'status': ['In Progress', 'Completed'],
                'min_budget': '50',
                'max_budget': '100',
            }
        )
        self.assertEqual(response.status_code, 200)
        projects = list(response.context['projects'])
        self.assertEqual([p.project_id for p in projects], ['P3', 'P4'])

    def test_blank_and_date_filters(self):
        response = self.client.get(
            reverse('project_list'),
            {
                'managing_unit': ['__blank__'],
                'start_date_from': '2024-01-01',
            }
        )
        self.assertEqual(response.status_code, 200)
        projects = list(response.context['projects'])
        self.assertEqual(len(projects), 1)
        self.assertEqual(projects[0].project_id, 'P4')

    def test_list_sorting_and_collapsed_filter_summary(self):
        response = self.client.get(
            reverse('project_list'),
            {'sort': 'total_budget', 'status': ['In Progress']},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            [project.project_id for project in response.context['projects']],
            ['P4', 'P3', 'P1'],
        )
        self.assertEqual(response.context['active_filter_count'], 1)
        self.assertContains(response, '已应用 1 类条件')
        self.assertContains(response, 'id="filter-section" style="display: none;"')
        self.assertContains(response, 'name="sort" value="total_budget"')
        self.assertNotContains(response, 'projectListView')

    def test_project_list_is_paginated_without_duplicate_card_view(self):
        for index in range(5, 36):
            Project.objects.create(
                project_id=f'P{index}', name=f'Project {index}', start_year=2024,
                status='In Progress', directory_path=f'P{index}',
            )

        first_page = self.client.get(reverse('project_list'))
        second_page = self.client.get(reverse('project_list'), {'page': 2})

        self.assertEqual(len(first_page.context['projects']), 30)
        self.assertEqual(len(second_page.context['projects']), 5)
        self.assertContains(first_page, '共 35 项')
        self.assertNotContains(first_page, 'id="card-view"')

    def test_invalid_sort_falls_back_to_year_descending(self):
        response = self.client.get(reverse('project_list'), {'sort': 'not-a-field'})

        self.assertEqual(response.context['sort'], '-start_year')
        self.assertEqual(
            [project.project_id for project in response.context['projects']],
            ['P2', 'P3', 'P4', 'P1'],
        )


class ProgressMonitorFilterTests(TestCase):
    def setUp(self):
        self.admin_user = get_user_model().objects.create_user('progress-admin', password='StrongPass!234', is_staff=True)
        self.client.force_login(self.admin_user)
        today = timezone.localdate()

        Project.objects.create(
            project_id='PM_OVERDUE',
            name='Overdue Project',
            ownership='GroupA',
            managing_unit='UnitA',
            level='LevelA',
            project_type='TypeA',
            role='Lead',
            start_year=today.year - 1,
            status='In Progress',
            start_date=today - timedelta(days=30),
            planned_end_date=today - timedelta(days=1),
            directory_path='PM_OVERDUE',
        )
        Project.objects.create(
            project_id='PM_MIDTERM',
            name='Midterm Project',
            ownership='GroupA',
            managing_unit='UnitA',
            level='LevelA',
            project_type='TypeA',
            role='Lead',
            start_year=today.year,
            status='In Progress',
            start_date=today - timedelta(days=20),
            planned_end_date=today + timedelta(days=20),
            directory_path='PM_MIDTERM',
        )
        Project.objects.create(
            project_id='PM_ONTRACK',
            name='OnTrack Project',
            ownership='GroupA',
            managing_unit='UnitA',
            level='LevelA',
            project_type='TypeA',
            role='Lead',
            start_year=today.year,
            status='In Progress',
            start_date=today - timedelta(days=10),
            planned_end_date=today + timedelta(days=40),
            directory_path='PM_ONTRACK',
        )

    def test_progress_status_filter_overdue(self):
        response = self.client.get(reverse('progress_monitor'), {'progress_status': 'overdue'})
        self.assertEqual(response.status_code, 200)

        rows = list(response.context['monitor_rows'])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['project'].project_id, 'PM_OVERDUE')
        self.assertEqual(rows[0]['status_key'], 'overdue')
        self.assertEqual(rows[0]['end_date'], timezone.localdate() - timedelta(days=1))
        self.assertEqual(response.context['selected_progress_status'], 'overdue')

    def test_extension_date_overrides_original_planned_end_date(self):
        today = timezone.localdate()
        extension_date = today + timedelta(days=30)
        Project.objects.create(
            project_id='PM_EXTENDED',
            name='Extended Project',
            ownership='西勘院',
            managing_unit='UnitA',
            level='省部级',
            project_type='应用研究',
            role='牵头',
            start_year=today.year,
            status='延期',
            start_date=today - timedelta(days=10),
            planned_end_date=today - timedelta(days=1),
            extension_date=extension_date,
            directory_path='PM_EXTENDED',
        )

        response = self.client.get(reverse('progress_monitor'), {'q': 'PM_EXTENDED'})

        self.assertEqual(response.status_code, 200)
        rows = list(response.context['monitor_rows'])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['end_date'], extension_date)
        self.assertEqual(rows[0]['days_remaining'], 30)
        self.assertEqual(rows[0]['status_key'], 'ontrack')

    def test_progress_status_filter_midterm(self):
        response = self.client.get(reverse('progress_monitor'), {'progress_status': 'midterm'})
        self.assertEqual(response.status_code, 200)

        rows = list(response.context['monitor_rows'])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['project'].project_id, 'PM_MIDTERM')
        self.assertEqual(rows[0]['status_key'], 'midterm')

    def test_progress_cards_prioritize_overdue_extended_then_midterm(self):
        today = timezone.localdate()
        priority_projects = [
            ('PM_PRIORITY_COMPLETE', '结题', today.year + 2),
            ('PM_PRIORITY_NOT_ESTABLISHED', '未立项', today.year + 1),
            ('PM_PRIORITY_ONGOING', '在研', today.year - 10),
            ('PM_PRIORITY_EXTENDED', '延期', today.year - 11),
        ]
        for project_id, status, start_year in priority_projects:
            Project.objects.create(
                project_id=project_id,
                name=project_id,
                start_year=start_year,
                status=status,
                start_date=today,
                planned_end_date=today + timedelta(days=100),
                directory_path=project_id,
            )

        response = self.client.get(reverse('progress_monitor'))

        self.assertEqual(response.status_code, 200)
        rows = list(response.context['monitor_rows'])
        project_ids = [row['project'].project_id for row in rows]
        self.assertEqual(
            project_ids[:3],
            ['PM_OVERDUE', 'PM_PRIORITY_EXTENDED', 'PM_MIDTERM'],
        )
        self.assertLess(
            project_ids.index('PM_PRIORITY_ONGOING'),
            project_ids.index('PM_PRIORITY_COMPLETE'),
        )

    def test_progress_monitor_is_paginated(self):
        today = timezone.localdate()
        for index in range(31):
            Project.objects.create(
                project_id=f'PM_EXTRA_{index:02d}', name=f'Extra {index}',
                start_year=today.year, status='In Progress',
                start_date=today, planned_end_date=today + timedelta(days=100),
                directory_path=f'PM_EXTRA_{index:02d}',
            )

        first_page = self.client.get(reverse('progress_monitor'))
        second_page = self.client.get(reverse('progress_monitor'), {'page': 2})

        self.assertEqual(len(first_page.context['monitor_rows']), 30)
        self.assertEqual(len(second_page.context['monitor_rows']), 4)
        self.assertContains(first_page, '共 34 项')


class QueryAssistantTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('assistant-reader', password='StrongPass!234')

        def create_project(project_id, name, start_year, **overrides):
            values = {
                'ownership': '西勘院',
                'managing_unit': '科技部',
                'level': '省部级',
                'project_type': '应用研究',
                'role': '牵头',
                'status': '在研',
                'start_date': date(start_year, 1, 1),
                'planned_end_date': date(2028, 12, 31),
                'directory_path': project_id,
            }
            values.update(overrides)
            return Project.objects.create(
                project_id=project_id,
                name=name,
                start_year=start_year,
                **values,
            )

        create_project(
            'ASSIST-2025', '省部级历史课题', 2025,
            project_lead='李四',
        )
        create_project(
            'ASSIST-2026-A', '张三负责的年度课题', 2026,
            project_lead='张三',
            planned_end_date=date(2026, 10, 31),
        )
        create_project(
            'ASSIST-2026-B', '张三联系的年度课题', 2026,
            project_lead='王五',
            contact_person='张三',
            planned_end_date=date(2026, 12, 31),
        )
        create_project(
            'ASSIST-EXTENDED-OUT', '已延期到下一年的课题', 2024,
            planned_end_date=date(2026, 6, 30),
            extension_date=date(2027, 6, 30),
        )
        create_project(
            'ASSIST-EXTENDED-IN', '延期至今年的课题', 2024,
            planned_end_date=date(2025, 6, 30),
            extension_date=date(2026, 9, 30),
        )
        create_project(
            'ASSIST-CLOSED', '今年已经结题的课题', 2024,
            status='结题',
            planned_end_date=date(2026, 3, 31),
            actual_completion_date=date(2026, 3, 20),
        )

    def test_finds_projects_for_lead_or_contact_person(self):
        result = answer_project_question('给我查一下张三有哪些课题', use_ai=False)

        self.assertTrue(result['ok'])
        self.assertEqual(result['kind'], 'project_list')
        self.assertEqual(result['total'], 2)
        self.assertEqual(
            {row['project'].project_id for row in result['rows']},
            {'ASSIST-2026-A', 'ASSIST-2026-B'},
        )
        self.assertEqual(
            {row['person_roles'] for row in result['rows']},
            {'负责人', '联系人'},
        )

    def test_lists_current_year_completion_using_extension_first(self):
        result = answer_project_question(
            '今年有哪些课题要结题',
            today=date(2026, 8, 21),
            use_ai=False,
        )

        self.assertTrue(result['ok'])
        self.assertEqual(result['total'], 3)
        project_ids = [row['project'].project_id for row in result['rows']]
        self.assertIn('ASSIST-EXTENDED-IN', project_ids)
        self.assertNotIn('ASSIST-EXTENDED-OUT', project_ids)
        self.assertNotIn('ASSIST-CLOSED', project_ids)

    def test_compares_two_and_four_digit_years_with_growth_rate(self):
        short_year_result = answer_project_question(
            '26年比25年省部级课题增长数量和比率',
            use_ai=False,
        )
        full_year_result = answer_project_question(
            '2026年比2025年省部级课题增长数量和比率',
            use_ai=False,
        )

        for result in (short_year_result, full_year_result):
            self.assertTrue(result['ok'])
            self.assertEqual(result['comparison']['target_count'], 2)
            self.assertEqual(result['comparison']['base_count'], 1)
            self.assertEqual(result['comparison']['difference'], 1)
            self.assertEqual(result['comparison']['growth_rate_text'], '100.0%')

    def test_ai_tools_support_search_grouping_and_comparison(self):
        search_result, project_ids = _tool_find_projects({
            'filters': {'person': '张三', 'person_role': 'any'},
            'sort_by': 'start_year',
            'limit': 10,
        })
        summary_result, _ = _tool_summarize_projects({
            'filters': {'level': ['省部级']},
            'group_by': 'start_year',
            'metric': 'count',
        })
        comparison_result, _ = _tool_compare_project_years({
            'years': [2026, 2025],
            'year_field': 'start_year',
            'filters': {'level': ['省部级'], 'start_year': 2026},
            'metric': 'count',
        })

        self.assertEqual(search_result['total_matches'], 2)
        self.assertEqual(set(project_ids), {'ASSIST-2026-A', 'ASSIST-2026-B'})
        groups = {row['group']: row['value'] for row in summary_result['groups']}
        self.assertEqual(groups[2026], 2)
        self.assertEqual(groups[2025], 1)
        self.assertEqual(comparison_result['values'][0]['value'], 2)
        self.assertEqual(comparison_result['values'][1]['value'], 1)
        self.assertEqual(comparison_result['first_vs_last']['change_rate_percent'], 100.0)

    def test_deepseek_tool_call_result_is_used_for_final_answer(self):
        class Config:
            @staticmethod
            def get_api_key():
                return 'sk-test-key-for-unit-test'

        tool_response = {
            'choices': [{
                'message': {
                    'role': 'assistant',
                    'content': None,
                    'tool_calls': [{
                        'id': 'call-1',
                        'type': 'function',
                        'function': {
                            'name': 'summarize_projects',
                            'arguments': json.dumps({
                                'filters': {'level': ['省部级']},
                                'group_by': 'start_year',
                                'metric': 'count',
                            }),
                        },
                    }],
                },
            }],
        }
        final_response = {
            'choices': [{
                'message': {
                    'role': 'assistant',
                    'content': '- **2026年省部级课题2项**，2025年1项。',
                },
            }],
        }

        with patch('core.ai_query_assistant._deepseek_config', return_value=Config()), \
                patch('core.ai_query_assistant._post_deepseek', side_effect=[tool_response, final_response]) as api_call:
            result = answer_with_deepseek(
                '省部级课题这两年有什么变化？',
                history=[{'role': 'user', 'content': '比较2026年和2025年'}],
                today=date(2026, 8, 21),
            )

        self.assertTrue(result['ok'])
        self.assertEqual(result['kind'], 'ai_answer')
        self.assertEqual(result['source'], 'DeepSeek AI')
        self.assertIn('2026年省部级课题2项', result['answer'])
        self.assertNotIn('**', result['answer'])
        self.assertTrue(result['answer'].startswith('• '))
        self.assertEqual(len(result['tool_summaries']), 1)
        self.assertEqual(api_call.call_count, 2)
        first_payload = api_call.call_args_list[0].args[1]
        self.assertEqual(first_payload['thinking'], {'type': 'disabled'})

    def test_readonly_user_can_use_assistant_widget_endpoint(self):
        request = RequestFactory().post(
            reverse('query_assistant'),
            data=json.dumps({'question': '张三有哪些课题', 'history': []}),
            content_type='application/json',
        )
        request.user = self.user
        middleware = AccessControlMiddleware(query_assistant_view)
        with patch('core.ai_query_assistant._deepseek_config', return_value=None):
            response = middleware(request)

        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        self.assertTrue(payload['ok'])
        self.assertIn('张三负责的年度课题', payload['html'])
        self.assertIn('张三联系的年度课题', payload['html'])
        self.assertIn('负责人或联系人', payload['html'])
        self.assertIn('history_content', payload)
        self.assertIn('query_assistant', AccessControlMiddleware.readonly_get_views)

    def test_assistant_endpoint_forwards_selected_local_model(self):
        request = RequestFactory().post(
            reverse('query_assistant'),
            data=json.dumps({
                'question': '今年有哪些课题要结题',
                'history': [],
                'service_name': 'local',
            }),
            content_type='application/json',
        )
        request.user = self.user
        with patch('core.views.answer_project_question', return_value={
            'ok': True,
            'kind': 'ai_answer',
            'answer': '测试回答',
            'source': '本地模型 / XKY-AI',
            'scope': '只读查询',
        }) as answer:
            response = query_assistant_view(request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(answer.call_args.kwargs['service_name'], 'local')

    def test_widget_lists_ready_local_model(self):
        APIConfig.objects.update_or_create(
            service_name='local',
            defaults={
                'api_key': '',
                'base_url': 'http://192.168.0.182:8000',
                'model_name': 'XKY-AI',
                'is_active': True,
                'test_success': True,
            },
        )
        request = RequestFactory().get(reverse('project_list'))
        request.user = self.user
        html = render_to_string('core/partials/navbar.html', request=request)

        self.assertIn('data-ai-model-select', html)
        self.assertIn('value="local"', html)
        self.assertIn('本地模型 / XKY-AI', html)

    def test_navigation_renders_floating_assistant_instead_of_nav_link(self):
        request = RequestFactory().get(reverse('project_list'))
        request.user = self.user
        html = render_to_string('core/partials/navbar.html', request=request)

        self.assertIn('data-ai-widget', html)
        self.assertIn('data-ai-launcher', html)
        self.assertIn('课题智能助手', html)
        self.assertNotIn('href="/assistant/" class="navbar-nav-link', html)


class ProjectCreationTests(TestCase):
    def test_create_project_page_renders_required_role_field(self):
        html = render_to_string('core/create_project.html', {'form': ProjectForm()})

        self.assertIn('name="role"', html)
        self.assertIn('class="container create-workspace"', html)
        self.assertIn('id="completion-bar"', html)
        self.assertIn('name="contact_person"', html)
        self.assertIn('name="research_content_manual"', html)
        self.assertIn('name="remarks"', html)
        self.assertIn('accept=".doc,.docx,.pdf"', html)
        self.assertIn('Word 或 PDF 任务书', html)
        self.assertNotIn('章节导航', html)

    def test_pdf_task_text_maps_core_project_fields(self):
        extracted = extract_task_pdf_text('''
            课题编号：PDF-2026-01
            课题名称：PDF任务书解析测试
            课题承担单位：西勘院技术中心
            课题负责人：张三
            课题联系人：李四
            课题起止年限：2026年1月1日 - 2027年12月31日
            外部专项经费：20
            院专项经费：10
            所属单位自筹经费：5
            主要研究内容：研究PDF任务书自动识别与字段预填。
        ''', source_file='任务书.pdf')

        self.assertEqual(extracted['basic_info']['课题编号'], 'PDF-2026-01')
        self.assertEqual(extracted['basic_info']['课题名称'], 'PDF任务书解析测试')
        self.assertEqual(extracted['budget_summary']['rows'][0]['专项经费'], 20)
        self.assertEqual(extracted['budget_summary']['rows'][0]['院专项经费'], 10)
        self.assertIn('PDF任务书自动识别', extracted['topic_info']['fields']['主要研究内容'])

    @patch('core.views.extract_task_pdf')
    def test_quick_import_endpoint_accepts_pdf(self, extract_pdf):
        extract_pdf.return_value = {
            'basic_info': {'课题编号': 'PDF-2026-02', '课题名称': 'PDF接口测试'},
            'topic_info': {'fields': {}},
            'budget_summary': {},
        }
        upload = SimpleUploadedFile('任务书.pdf', b'%PDF-1.7 test', content_type='application/pdf')
        request = RequestFactory().post(reverse('extract_task_docx'), {'document': upload})

        response = extract_task_docx_view(request)
        payload = json.loads(response.content)

        self.assertEqual(response.status_code, 200)
        self.assertTrue(payload['success'])
        self.assertEqual(payload['data']['project_id'], 'PDF-2026-02')
        extract_pdf.assert_called_once()

    def test_scanned_pdf_text_has_clear_ocr_message(self):
        with self.assertRaisesRegex(ValueError, 'OCR'):
            extract_task_pdf_text('   \n\t  ', source_file='扫描任务书.pdf')

    def test_admin_navigation_exposes_create_project_action(self):
        html = render_to_string('core/partials/navbar.html', {'is_system_admin': True})

        self.assertIn('class="navbar-create-btn', html)
        self.assertIn('新建课题', html)

    def test_project_level_choices_include_city_level(self):
        form = ProjectForm()
        level_values = [value for value, _ in form.fields['level'].choices if value]

        self.assertEqual(level_values, ['国家级', '省部级', '地市级', '公司级'])

    def test_project_status_choices_are_limited_and_legacy_values_normalize(self):
        form = ProjectForm()
        status_values = [value for value, _ in form.fields['status'].choices if value]

        self.assertEqual(status_values, ['未立项', '在研', '延期', '结题', '终止'])
        self.assertEqual(Project.normalize_status('未立项'), '未立项')
        self.assertEqual(Project.normalize_status('申报'), '未立项')
        self.assertEqual(Project.normalize_status('立项'), '在研')
        self.assertEqual(Project.normalize_status('已立项'), '在研')
        self.assertEqual(Project.normalize_status('完成'), '结题')
        self.assertEqual(Project.normalize_status('未知状态'), '')

    def test_import_strips_text_fields_before_creating_directory(self):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.append([
            '课题编号', '课题名称', '课题归属', '归口单位', '课题级别',
            '课题类型', '参与角色', '开始年份', '课题状态',
        ])
        worksheet.append([
            'IMPORT-TRIM-1', '导入课题末尾空格 ', '西勘院', '测试单位', '公司级',
            '应用研究', '牵头', 2026, '在研',
        ])
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        upload = SimpleUploadedFile(
            'projects.xlsx',
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        with tempfile.TemporaryDirectory() as projects_root:
            with override_settings(PROJECTS_ROOT=Path(projects_root)):
                request = RequestFactory().post(reverse('import_from_excel'), {'excel_file': upload})
                request.session = {}
                request._messages = FallbackStorage(request)
                response = import_from_excel_view(request)

                self.assertEqual(response.status_code, 302)
                self.assertEqual(response.url, reverse('project_list'))
                project = Project.objects.get(project_id='IMPORT-TRIM-1')
                self.assertEqual(project.name, '导入课题末尾空格')
                self.assertEqual(Path(project.directory_path).name, '2026-在研-IMPORT-TRIM-1-导入课题末尾空格')
                self.assertTrue(Path(project.directory_path, '01_申报').is_dir())


class BackupScheduleTests(TestCase):
    weekly_task_payload = (
        '{"task_name":"KetiBackupWeekly","state":"Ready","enabled":true,'
        '"start_boundary":"2026-02-06T02:00:00","time":"02:00",'
        '"trigger_type":"Weekly","interval_days":7,'
        '"last_run_time":"2026-07-26 02:00:00",'
        '"next_run_time":"2026-08-02 02:00:00","last_result":0}'
    )
    daily_task_payload = (
        '{"task_name":"KetiBackupWeekly","state":"Ready","enabled":true,'
        '"start_boundary":"2026-07-28T03:30:00","time":"03:30",'
        '"trigger_type":"Daily","interval_days":3,'
        '"last_run_time":"2026-07-26 02:00:00",'
        '"next_run_time":"2026-07-31 03:30:00","last_result":0}'
    )

    def setUp(self):
        self.admin_user = get_user_model().objects.create_user('backup-admin', password='StrongPass!234', is_staff=True)
        self.client.force_login(self.admin_user)

    @override_settings(BACKUP_TASK_NAME=r'\KetiBackupWeekly', BACKUP_SCHEDULE_READ_ONLY=False)
    @patch('core.backup_schedule._run_powershell')
    def test_reads_weekly_backup_schedule(self, run_powershell):
        run_powershell.return_value = self.weekly_task_payload

        schedule = backup_schedule.get_backup_schedule()

        read_script = run_powershell.call_args.args[0]
        self.assertIn('MSFT_TaskDailyTrigger', read_script)
        self.assertIn('MSFT_TaskWeeklyTrigger', read_script)
        self.assertIn('WeeksInterval * 7', read_script)
        self.assertTrue(schedule['available'])
        self.assertEqual(schedule['trigger_type'], 'Weekly')
        self.assertEqual(schedule['interval_days'], 7)
        self.assertEqual(schedule['time'], '02:00')
        self.assertTrue(schedule['last_result_success'])

    @override_settings(BACKUP_TASK_NAME=r'\KetiBackupWeekly', BACKUP_SCHEDULE_READ_ONLY=False)
    @patch('core.backup_schedule._run_powershell')
    def test_updates_only_daily_interval_trigger(self, run_powershell):
        run_powershell.side_effect = ['', self.daily_task_payload]

        schedule = backup_schedule.update_backup_schedule('3', '03:30', True)

        update_script = run_powershell.call_args_list[0].args[0]
        self.assertIn('Set-ScheduledTask', update_script)
        self.assertIn('New-ScheduledTaskTrigger -Daily -DaysInterval 3', update_script)
        self.assertIn('AddHours(3).AddMinutes(30)', update_script)
        self.assertIn('Enable-ScheduledTask', update_script)
        self.assertNotIn('DaysOfWeek', update_script)
        self.assertNotIn('backup_weekly.ps1', update_script)
        self.assertTrue(schedule['available'])
        self.assertEqual(schedule['interval_days'], 3)

    @override_settings(BACKUP_TASK_NAME=r'\KetiBackupWeekly', BACKUP_SCHEDULE_READ_ONLY=False)
    @patch('core.backup_schedule._run_powershell')
    def test_can_disable_automatic_backup_without_changing_task_action(self, run_powershell):
        disabled_payload = self.daily_task_payload.replace('"enabled":true', '"enabled":false')
        run_powershell.side_effect = ['', disabled_payload]

        schedule = backup_schedule.update_backup_schedule('15', '02:00', False)

        update_script = run_powershell.call_args_list[0].args[0]
        self.assertIn('Disable-ScheduledTask', update_script)
        self.assertNotIn('backup_weekly.ps1', update_script)
        self.assertFalse(schedule['enabled'])

    @override_settings(BACKUP_SCHEDULE_READ_ONLY=True)
    @patch('core.backup_schedule._run_powershell')
    def test_development_read_only_mode_blocks_updates(self, run_powershell):
        schedule = backup_schedule.get_backup_schedule()
        self.assertTrue(schedule['available'])
        self.assertTrue(schedule['read_only'])
        self.assertTrue(schedule['preview'])
        self.assertEqual(schedule['interval_days'], 7)
        with self.assertRaisesMessage(backup_schedule.BackupScheduleError, '开发环境为只读模式'):
            backup_schedule.update_backup_schedule('7', '02:00', True)
        run_powershell.assert_not_called()

    @override_settings(BACKUP_SCHEDULE_READ_ONLY=False)
    @patch('core.backup_schedule._run_powershell')
    def test_invalid_schedule_values_are_rejected_before_powershell(self, run_powershell):
        with self.assertRaisesMessage(backup_schedule.BackupScheduleError, '必须是整数天数'):
            backup_schedule.update_backup_schedule('three', '02:00', True)
        with self.assertRaisesMessage(backup_schedule.BackupScheduleError, '1 到 3650'):
            backup_schedule.update_backup_schedule('0', '02:00', True)
        with self.assertRaisesMessage(backup_schedule.BackupScheduleError, '格式为 HH:MM'):
            backup_schedule.update_backup_schedule('3', '25:99', True)
        run_powershell.assert_not_called()

    def test_settings_page_shows_adjust_backup_button(self):
        schedule = {
            'available': True,
            'read_only': False,
            'task_name': 'KetiBackupWeekly',
            'state': 'Ready',
            'enabled': True,
            'trigger_type': 'Weekly',
            'interval_days': 7,
            'time': '02:00',
            'last_run_time': '2026-07-26 02:00:00',
            'next_run_time': '2026-08-02 02:00:00',
            'last_result': 0,
            'last_result_success': True,
        }

        html = render_to_string(
            'core/partials/backup_schedule_settings.html',
            {
                'backup_schedule': schedule,
            },
        )

        self.assertIn('调整备份计划', html)
        self.assertIn('name="backup_interval_days"', html)
        self.assertIn('每隔 7 天，于 02:00 执行', html)

    @patch('core.views.backup_scheduler.get_backup_schedule')
    def test_backup_schedule_status_loads_from_async_endpoint(self, get_backup_schedule):
        get_backup_schedule.return_value = {
            'available': True, 'read_only': False, 'task_name': 'KetiBackupWeekly',
            'enabled': True, 'interval_days': 3, 'time': '04:15',
            'last_result_success': True,
        }

        settings_response = self.client.get(reverse('settings'))
        status_response = self.client.get(reverse('backup_schedule_status'))

        self.assertEqual(settings_response.status_code, 200)
        self.assertContains(settings_response, 'data-status-url=')
        self.assertEqual(status_response.status_code, 200)
        self.assertContains(status_response, '每隔 3 天，于 04:15 执行')
        get_backup_schedule.assert_called_once_with()

    @patch('core.views.backup_scheduler.update_backup_schedule')
    def test_settings_page_updates_backup_schedule(self, update_schedule):
        update_schedule.return_value = {'available': True}

        response = self.client.post(
            reverse('settings'),
            {
                'action': 'save_backup_schedule',
                'backup_interval_days': '3',
                'backup_time': '04:15',
                'backup_enabled': 'on',
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('settings'))
        update_schedule.assert_called_once_with('3', '04:15', True)


class MetricCatalogConfigurationTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_user(
            'metric-config-admin',
            password='StrongPass!234',
            is_staff=True,
        )
        self.client.force_login(self.admin)
        self.project = Project.objects.create(
            project_id='METRIC-CONFIG-1',
            name='指标配置测试课题',
            ownership='西勘院',
            managing_unit='测试单位',
            level='公司级',
            project_type='应用研究',
            role='牵头',
            start_year=2026,
            status='在研',
            directory_path='METRIC-CONFIG-1',
        )

    def test_settings_page_lists_seeded_metric_catalog(self):
        response = self.client.get(reverse('settings'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="metric-settings"')
        self.assertContains(response, '指标类型管理')
        self.assertContains(response, '知识产权类')
        self.assertContains(response, '软件著作权')
        self.assertContains(response, '图纸')

    def test_custom_indicator_from_settings_is_available_when_creating_metric(self):
        category_response = self.client.post(reverse('settings'), {
            'action': 'create_metric_category',
            'category_name': '人才培养类',
            'sort_order': '50',
            'is_active': 'on',
        })
        self.assertEqual(category_response.status_code, 302)
        category = MetricsCategory.objects.get(name='人才培养类')

        indicator_response = self.client.post(reverse('settings'), {
            'action': 'create_metric_indicator',
            'category_id': str(category.id),
            'indicator_name': '技术骨干培养',
            'unit': '人',
            'assessment_method': '培养名单及证明',
            'sort_order': '10',
            'is_active': 'on',
        })
        self.assertEqual(indicator_response.status_code, 302)

        catalog_item = MetricsItem.get_catalog_item('技术骨干培养')
        self.assertEqual(catalog_item['category'], category.code)
        self.assertEqual(catalog_item['unit'], '人')
        self.assertEqual(catalog_item['assessment_method'], '培养名单及证明')

        create_response = self.client.post(
            reverse('create_metrics_item', args=[self.project.project_id]),
            {'item_name': '技术骨干培养', 'target_value': '3'},
        )
        self.assertEqual(create_response.status_code, 302)
        metric = MetricsItem.objects.get(item_name='技术骨干培养')
        self.assertEqual(metric.category, category.code)
        self.assertEqual(metric.target_value, '3人')
        self.assertEqual(metric.assessment_method, '培养名单及证明')

    def test_disabling_indicator_hides_it_but_keeps_existing_metrics(self):
        indicator = MetricIndicatorDefinition.objects.get(name='软件著作权')
        analysis = ProjectAnalysis.objects.create(
            project=self.project,
            analysis_type='output_metrics',
            file_name='手动录入',
            analysis_result='指标',
        )
        metric = MetricsItem.objects.create(
            analysis=analysis,
            category=indicator.category.code,
            item_name=indicator.name,
            target_value='1项',
        )

        response = self.client.post(reverse('settings'), {
            'action': 'update_metric_indicator',
            'indicator_id': str(indicator.id),
            'category_id': str(indicator.category_id),
            'indicator_name': indicator.name,
            'unit': indicator.unit,
            'assessment_method': indicator.assessment_method,
            'sort_order': str(indicator.sort_order),
        })

        self.assertEqual(response.status_code, 302)
        indicator.refresh_from_db()
        self.assertFalse(indicator.is_active)
        self.assertIsNone(MetricsItem.get_catalog_item('软件著作权'))
        self.assertTrue(MetricsItem.objects.filter(pk=metric.pk).exists())

    def test_category_rename_updates_existing_metric_group_label(self):
        category = MetricsCategory.objects.get(code='ip')
        response = self.client.post(reverse('settings'), {
            'action': 'update_metric_category',
            'category_id': str(category.id),
            'category_name': '知识产权成果类',
            'sort_order': str(category.sort_order),
            'is_active': 'on',
        })

        self.assertEqual(response.status_code, 302)
        self.assertEqual(MetricsItem.get_category_label_map()['ip'], '知识产权成果类')


class AccessControlTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user('admin-user', password='StrongPass!234', is_staff=True)
        self.readonly = User.objects.create_user('readonly-user', password='StrongPass!234')
        self.project = Project.objects.create(
            project_id='ACCESS-1',
            name='权限测试课题',
            ownership='西勘院',
            managing_unit='测试单位',
            level='公司级',
            project_type='应用研究',
            role='牵头',
            start_year=2026,
            status='在研',
            total_budget=Decimal('99.90'),
            external_funding=Decimal('12.30'),
            institute_funding=Decimal('4.50'),
            directory_path='ACCESS-1',
        )

    def test_anonymous_user_is_redirected_to_login(self):
        response = self.client.get(reverse('project_list'))

        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith(reverse('login')))

    def test_readonly_user_can_view_content_but_not_admin_pages_or_writes(self):
        self.client.force_login(self.readonly)

        self.assertEqual(self.client.get(reverse('project_list')).status_code, 200)
        self.assertEqual(self.client.get(reverse('statistics')).status_code, 200)
        self.assertEqual(self.client.get(reverse('progress_monitor')).status_code, 200)
        self.assertEqual(self.client.get(reverse('settings')).status_code, 403)
        self.assertEqual(self.client.post(reverse('create_project'), {}).status_code, 403)

    def test_readonly_page_hides_admin_controls(self):
        self.client.force_login(self.readonly)

        response = self.client.get(reverse('project_list'))

        self.assertNotContains(response, '从Excel导入')
        self.assertNotContains(response, '新建课题')
        self.assertNotContains(response, 'API配置')
        self.assertContains(response, '只读')
        self.assertContains(response, '外部专项')
        self.assertContains(response, '12.3')
        self.assertContains(response, '院专项')
        self.assertContains(response, '4.5')
        self.assertNotContains(response, "onclick=\"confirmDelete('ACCESS-1'")
        self.assertContains(response, 'class="navbar-toggle"')

    def test_admin_navigation_uses_grouped_business_and_configuration_entries(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse('project_list'))

        self.assertContains(response, '>课题管理<')
        self.assertContains(response, '>监控中心<')
        self.assertContains(response, '>系统配置<')
        self.assertNotContains(response, '>API配置<')
        self.assertNotContains(response, 'title="用户管理"')

    def test_configuration_pages_share_module_navigation(self):
        self.client.force_login(self.admin)

        for route_name in ('settings', 'api_config', 'user_management'):
            response = self.client.get(reverse(route_name))
            self.assertEqual(response.status_code, 200)
            self.assertContains(response, 'aria-label="系统配置模块"')
            self.assertContains(response, 'AI 模型配置')
            self.assertContains(response, '用户与权限')

    def test_project_table_uses_separate_external_and_institute_funding_columns(self):
        self.client.force_login(self.readonly)

        response = self.client.get(reverse('project_list'), {'sort': 'external_funding'})
        html = response.content.decode('utf-8')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['sort'], 'external_funding')
        self.assertContains(response, '外部专项')
        self.assertContains(response, '<th scope="col" class="col-budget">院专项</th>', html=True)
        self.assertIn('sort=-external_funding', html)
        self.assertRegex(
            html,
            r'<td class="col-budget">\s*<strong>12\.3</strong>\s*</td>\s*'
            r'<td class="col-budget">\s*<strong>4\.5</strong>',
        )
        self.assertNotRegex(
            html,
            r'<td class="col-budget">\s*<strong>(?:12\.3|4\.5)万</strong>',
        )
        self.assertNotRegex(
            html,
            r'<div class="card-budget-item">\s*(?:外部专项|院专项)\s*'
            r'<strong>(?:12\.3|4\.5)万(?:元)?</strong>',
        )
        self.assertNotIn('>专项预算</th>', html)

    def test_login_is_chinese_and_uses_local_icon_assets(self):
        response = self.client.get(reverse('login'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<label for="id_username">用户名</label>', html=True)
        self.assertContains(response, '<label for="id_password">密码</label>', html=True)
        self.assertContains(response, 'core/vendor/fontawesome/css/all.min.css')
        self.assertNotContains(response, 'cdnjs.cloudflare.com')

    def test_admin_table_view_keeps_delete_action(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse('project_list'))

        self.assertContains(response, 'title="删除课题"')
        self.assertContains(response, 'onclick="confirmDelete(')
        self.assertNotContains(response, 'id="card-view"')

    def test_user_management_uses_shared_page_structure(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse('user_management'))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'core/base.html')
        self.assertTemplateUsed(response, 'core/partials/page_header.html')
        self.assertContains(response, 'class="skip-link"')
        self.assertContains(response, 'id="create-username"')

    def test_api_config_uses_balanced_dashboard_layout(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse('api_config'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="api-setup-grid"')
        self.assertContains(response, 'class="content-section api-config-list"')
        self.assertContains(response, 'class="content-section api-instructions"')
        self.assertContains(response, 'core/css/navigation.css')
        self.assertContains(response, '?v=3.0')

    def test_readonly_preview_is_allowed_and_download_uses_global_switch(self):
        self.client.force_login(self.readonly)
        with tempfile.TemporaryDirectory() as config_root, tempfile.TemporaryDirectory() as projects_root:
            config_root_path = Path(config_root)
            project_root_path = Path(projects_root)
            folder = project_root_path / '2026-在研-ACCESS-1-权限测试课题'
            target_folder = folder / '01_申报'
            target_folder.mkdir(parents=True)
            (target_folder / '材料.txt').write_text('test', encoding='utf-8')
            self.project.directory_path = str(folder)
            self.project.save(update_fields=['directory_path'])
            config_file = config_root_path / 'network_config.json'
            config_file.write_text('{"readonly_can_download": false}', encoding='utf-8')

            with override_settings(BASE_DIR=config_root_path, PROJECTS_ROOT=project_root_path):
                detail_response = self.client.get(reverse('project_detail', args=[self.project.project_id]))
                self.assertEqual(detail_response.status_code, 200)
                self.assertContains(detail_response, 'class="page-header detail-page-header"')
                self.assertContains(detail_response, 'core/css/project-detail-header.css')
                self.assertContains(detail_response, 'class="page-title detail-project-title"')
                self.assertNotContains(detail_response, '<h3>编辑信息</h3>', html=True)
                self.assertNotContains(detail_response, '上传到 <span id="upload-target-name">根目录</span>', html=True)
                self.assertFalse(detail_response.context['can_download_files'])
                preview_url = reverse('file_action', args=[self.project.project_id, 'preview']) + '?path=01_申报/材料.txt'
                download_url = reverse('file_action', args=[self.project.project_id, 'download']) + '?path=01_申报/材料.txt'
                preview_response = self.client.get(preview_url)
                self.assertEqual(preview_response.status_code, 200)
                preview_response.close()
                self.assertEqual(self.client.get(download_url).status_code, 403)

                config_file.write_text('{"readonly_can_download": true}', encoding='utf-8')
                download_response = self.client.get(download_url)
                self.assertEqual(download_response.status_code, 200)
                download_response.close()

    def test_admin_can_create_readonly_user_and_operation_is_logged(self):
        self.client.force_login(self.admin)

        response = self.client.post(reverse('user_management'), {
            'action': 'create',
            'username': 'new-reader',
            'display_name': '只读测试用户',
            'role': 'readonly',
            'password1': 'AnotherPass!234',
            'password2': 'AnotherPass!234',
        })

        self.assertEqual(response.status_code, 302)
        created_user = get_user_model().objects.get(username='new-reader')
        self.assertFalse(created_user.is_staff)
        self.assertTrue(created_user.is_active)
        self.assertTrue(OperationLog.objects.filter(user=self.admin, action_name='user_management').exists())


    def test_admin_detail_has_accessible_editing_structure(self):
        self.client.force_login(self.admin)
        with tempfile.TemporaryDirectory() as projects_root:
            self.project.directory_path = str(Path(projects_root) / 'ACCESS-1')
            self.project.save(update_fields=['directory_path'])
            with override_settings(PROJECTS_ROOT=Path(projects_root)):
                response = self.client.get(reverse('project_detail', args=[self.project.project_id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'role="tablist"')
        self.assertContains(response, 'role="tab"', count=4)
        self.assertContains(response, 'id="project-edit-form"')
        self.assertContains(response, 'for="id_start_year"')
        self.assertContains(response, 'class="skip-link"')
        self.assertContains(response, 'beforeunload')

class LazyFileTreeTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_user(
            'file-tree-admin',
            password='StrongPass!234',
            is_staff=True,
        )
        self.client.force_login(self.admin)

    def test_project_detail_and_file_tree_api_load_one_level_at_a_time(self):
        with tempfile.TemporaryDirectory() as projects_root:
            projects_root_path = Path(projects_root)
            project_folder = projects_root_path / '2026-在研-LAZY-1-大型目录测试'
            deep_folder = project_folder / '08_其他' / '软件源码' / 'venv' / 'Lib' / 'site-packages'
            deep_folder.mkdir(parents=True)
            (deep_folder / 'large-package.py').write_text('print("ok")', encoding='utf-8')
            project = Project.objects.create(
                project_id='LAZY-1',
                name='大型目录测试',
                ownership='西勘院',
                managing_unit='测试单位',
                level='公司级',
                project_type='应用研究',
                role='牵头',
                start_year=2026,
                status='在研',
                directory_path=str(project_folder),
            )

            with override_settings(PROJECTS_ROOT=projects_root_path):
                with patch('core.views.get_directory_tree', side_effect=AssertionError('不应递归扫描')):
                    detail_response = self.client.get(reverse('project_detail', args=[project.project_id]))
                self.assertEqual(detail_response.status_code, 200)
                root_nodes = detail_response.context['file_tree']
                other_node = next(node for node in root_nodes if node['name'] == '08_其他')
                self.assertTrue(other_node['has_children'])
                self.assertEqual(other_node['children'], [])

                root_response = self.client.get(reverse('get_file_tree', args=[project.project_id]))
                self.assertEqual(root_response.status_code, 200)
                root_names = [node['name'] for node in root_response.json()['file_tree']]
                self.assertIn('08_其他', root_names)
                self.assertNotIn('软件源码', root_names)

                child_response = self.client.get(
                    reverse('get_file_tree', args=[project.project_id]),
                    {'path': '08_其他'},
                )
                self.assertEqual(child_response.status_code, 200)
                child_nodes = child_response.json()['file_tree']
                self.assertEqual([node['name'] for node in child_nodes], ['软件源码'])
                self.assertTrue(child_nodes[0]['has_children'])
                self.assertEqual(child_nodes[0]['children'], [])

                invalid_response = self.client.get(
                    reverse('get_file_tree', args=[project.project_id]),
                    {'path': '..'},
                )
                self.assertEqual(invalid_response.status_code, 400)

    def test_web_file_trial_renders_expandable_tree_with_files(self):
        with tempfile.TemporaryDirectory() as projects_root:
            projects_root_path = Path(projects_root)
            project_folder = projects_root_path / '2026-在研-TREE-1-目录展开测试'
            material_folder = project_folder / '06_结题' / '验收材料'
            material_folder.mkdir(parents=True)
            (material_folder / '结题报告.pdf').write_bytes(b'%PDF-1.4 test')
            project = Project.objects.create(
                project_id='TREE-1',
                name='目录展开测试',
                ownership='西勘院',
                managing_unit='测试单位',
                level='公司级',
                project_type='应用研究',
                role='牵头',
                start_year=2026,
                status='在研',
                directory_path=str(project_folder),
            )

            with override_settings(PROJECTS_ROOT=projects_root_path):
                response = self.client.get(reverse('file_manager_trial', args=[project.project_id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'role="tree"')
        self.assertContains(response, 'async function loadFolderChildren')
        self.assertContains(response, 'toggleFolderNode(node)')
        self.assertContains(response, 'className = "tree-file-btn"')
        self.assertContains(response, '点击箭头逐级展开目录')


class DirectoryReconcileTests(TestCase):
    def test_dry_run_then_apply_preserves_conflicts_and_quarantines_old_folder(self):
        from .views import _get_project_folder_name

        with tempfile.TemporaryDirectory() as temp_root:
            temp_root_path = Path(temp_root)
            projects_root = temp_root_path / 'projects'
            quarantine_root = temp_root_path / 'quarantine'
            projects_root.mkdir()
            project = Project.objects.create(
                project_id='DUP-1',
                name='Current Project',
                ownership='西勘院',
                managing_unit='测试单位',
                level='公司级',
                project_type='应用研究',
                role='牵头',
                start_year=2026,
                status='在研',
                directory_path='',
            )
            expected = projects_root / _get_project_folder_name(project)
            duplicate = projects_root / '2025-延期-DUP-1-Old Project'
            expected.mkdir()
            duplicate.mkdir()
            (expected / 'same.txt').write_text('same', encoding='utf-8')
            (duplicate / 'same.txt').write_text('same', encoding='utf-8')
            (expected / 'conflict.txt').write_text('current', encoding='utf-8')
            (duplicate / 'conflict.txt').write_text('old', encoding='utf-8')
            (duplicate / 'only-old.txt').write_text('copy me', encoding='utf-8')

            dry_report = temp_root_path / 'dry.json'
            with override_settings(PROJECTS_ROOT=projects_root):
                call_command(
                    'reconcile_project_directories',
                    report=str(dry_report),
                    stdout=StringIO(),
                )
            self.assertTrue(duplicate.exists())
            self.assertFalse((expected / 'only-old.txt').exists())
            dry_data = __import__('json').loads(dry_report.read_text(encoding='utf-8'))
            self.assertEqual(dry_data['summary']['duplicate_directories'], 1)
            self.assertEqual(dry_data['summary']['errors'], 0)

            apply_report = temp_root_path / 'apply.json'
            with override_settings(PROJECTS_ROOT=projects_root):
                call_command(
                    'reconcile_project_directories',
                    apply=True,
                    quarantine_root=str(quarantine_root),
                    report=str(apply_report),
                    stdout=StringIO(),
                )

            apply_data = __import__('json').loads(apply_report.read_text(encoding='utf-8'))
            self.assertEqual(apply_data['summary']['errors'], 0)
            self.assertEqual(apply_data['summary']['quarantined_directories'], 1)
            self.assertFalse(duplicate.exists())
            self.assertEqual((expected / 'only-old.txt').read_text(encoding='utf-8'), 'copy me')
            self.assertEqual((expected / 'conflict.txt').read_text(encoding='utf-8'), 'current')
            conflict_copies = list(expected.glob('conflict__from_old_folder_*.txt'))
            self.assertEqual(len(conflict_copies), 1)
            self.assertEqual(conflict_copies[0].read_text(encoding='utf-8'), 'old')
            quarantined = list(quarantine_root.glob('*/*'))
            self.assertEqual(len(quarantined), 1)
            self.assertTrue((quarantined[0] / 'only-old.txt').exists())


class AnalysisTrackingTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_user(
            'analysis-admin',
            password='StrongPass!234',
            is_staff=True,
        )
        self.client.force_login(self.admin)
        self.project = Project.objects.create(
            project_id='ANALYSIS-1',
            name='任务书指标跟踪测试',
            ownership='西勘院',
            managing_unit='测试单位',
            level='公司级',
            project_type='应用研究',
            role='牵头',
            start_year=2026,
            status='在研',
            directory_path='ANALYSIS-1',
        )

    def test_manual_research_edit_replaces_structured_result(self):
        analysis = ProjectAnalysis.objects.create(
            project=self.project,
            analysis_type='research_content',
            file_name='任务书.docx',
            analysis_result='原研究目标\n\n原研究内容一\n\n原研究内容二',
            structured_data={
                'research_goal': '原研究目标',
                'research_sections': [
                    {'title': '内容一', 'description': '原研究内容一', 'subitems': []},
                    {'title': '内容二', 'description': '原研究内容二', 'subitems': []},
                ],
            },
        )

        response = self.client.post(
            reverse('edit_analysis', args=[self.project.project_id, 'research_content']),
            {'action': 'save', 'analysis_result': '仅保留修改后的研究内容'},
        )

        self.assertEqual(response.status_code, 302)
        analysis.refresh_from_db()
        self.project.refresh_from_db()
        self.assertEqual(analysis.analysis_result, '仅保留修改后的研究内容')
        self.assertFalse(any(analysis.structured_data.values()))
        self.assertEqual(self.project.research_content, '仅保留修改后的研究内容')
        self.assertNotIn('原研究内容一', analysis.analysis_result)

    def test_detail_offers_ready_model_for_both_ai_analysis_forms(self):
        APIConfig.objects.update_or_create(
            service_name='local',
            defaults={
                'api_key': '',
                'base_url': 'http://192.168.0.182:8000',
                'model_name': 'XKY-AI',
                'is_active': True,
                'test_success': True,
            },
        )
        with tempfile.TemporaryDirectory() as projects_root:
            with override_settings(PROJECTS_ROOT=Path(projects_root)):
                response = self.client.get(reverse('project_detail', args=[self.project.project_id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="ai_service"', count=2)
        self.assertContains(response, '本地模型 / XKY-AI')

    def test_document_analysis_forwards_selected_local_model(self):
        document = SimpleUploadedFile('task.txt', '课题研究目标和课题研究内容测试文本。'.encode('utf-8'))
        with patch('core.ai_analysis.ai_service.analyze_document', return_value={
            'success': False,
            'error': '测试停止',
        }) as analyze:
            response = self.client.post(
                reverse('analyze_content', args=[self.project.project_id, 'research_content']),
                {'files': document, 'ai_service': 'local'},
            )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(analyze.call_args.kwargs['service_name'], 'local')

    def test_research_analysis_can_be_deleted_and_project_content_is_cleared(self):
        ProjectAnalysis.objects.create(
            project=self.project,
            analysis_type='research_content',
            file_name='任务书.docx',
            analysis_result='待删除的研究内容',
            structured_data={'research_goal': '待删除的研究内容'},
        )

        response = self.client.post(
            reverse('edit_analysis', args=[self.project.project_id, 'research_content']),
            {'action': 'delete', 'analysis_result': '待删除的研究内容'},
        )

        self.assertRedirects(
            response,
            f"{reverse('project_detail', args=[self.project.project_id])}?tab=content-analysis",
            fetch_redirect_response=False,
        )
        self.assertFalse(
            ProjectAnalysis.objects.filter(
                project=self.project,
                analysis_type='research_content',
            ).exists()
        )
        self.project.refresh_from_db()
        self.assertEqual(self.project.research_content, '')

    def test_clearing_research_editor_deletes_stale_structured_result(self):
        ProjectAnalysis.objects.create(
            project=self.project,
            analysis_type='research_content',
            file_name='任务书.docx',
            analysis_result='旧内容',
            structured_data={'research_goal': '旧内容'},
        )

        self.client.post(
            reverse('edit_analysis', args=[self.project.project_id, 'research_content']),
            {'action': 'save', 'analysis_result': '   '},
        )

        self.assertFalse(
            ProjectAnalysis.objects.filter(
                project=self.project,
                analysis_type='research_content',
            ).exists()
        )

    def test_research_editor_renders_explicit_delete_button(self):
        ProjectAnalysis.objects.create(
            project=self.project,
            analysis_type='research_content',
            file_name='任务书.docx',
            analysis_result='可删除内容',
        )
        with tempfile.TemporaryDirectory() as projects_root:
            from .views import _get_project_folder_name

            expected_folder = Path(projects_root) / _get_project_folder_name(self.project)
            expected_folder.mkdir(parents=True)
            self.project.directory_path = str(expected_folder)
            self.project.save(update_fields=['directory_path'])
            with override_settings(PROJECTS_ROOT=Path(projects_root)):
                response = self.client.get(
                    reverse('project_detail', args=[self.project.project_id]),
                    {'tab': 'content-analysis'},
                )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'value="delete"')
        self.assertContains(response, '删除分析结果')

    def test_taskbook_json_preserves_original_indicator_descriptions(self):
        raw = '''
        {
          "summary": "按任务书表格提取",
          "metrics": [
            {"category": "deliverable", "indicator_description": "城市更新视角下老旧社区改造的策略研究论文", "target_value": "1", "assessment_method": "发表", "source_section": "考核指标", "source_page": "7"},
            {"category": "benefit", "indicator_description": "侠客岛城市更新项目", "target_value": "100万元", "planned_period": "2022年-2023年", "source_section": "主要经济、社会指标", "source_page": "7"},
            {"category": "milestone", "indicator_description": "完成资料收集", "planned_period": "2022.04-2022.05", "source_section": "课题进度计划", "source_page": "8"}
          ]
        }
        '''

        display, structured, metrics = parse_ai_analysis(raw, 'output_metrics')

        self.assertEqual(metrics[0]['item_name'], '城市更新视角下老旧社区改造的策略研究论文')
        self.assertEqual(metrics[0]['assessment_method'], '发表')
        self.assertEqual(metrics[1]['category'], 'benefit')
        self.assertEqual(metrics[2]['category'], 'milestone')
        self.assertIn('城市更新视角下老旧社区改造的策略研究论文', display)
        self.assertEqual(len(structured['metrics']), 3)

    def test_ai_prompt_requires_taskbook_wording_instead_of_fixed_templates(self):
        service = AIAnalysisService()
        prompt = service._build_prompt(
            {
                'filename': '常用课题任务书.pdf',
                'extension': 'pdf',
                'content': '[[任务书第7页]]\n二、预期成果及考核指标\n考核指标 数量 考核方式',
            },
            'output_metrics',
        )

        self.assertIn('完整原文', prompt)
        self.assertIn('主要经济、社会指标', prompt)
        self.assertIn('课题进度计划', prompt)
        self.assertNotIn('发明专利：[数量]项', prompt)

    def test_legacy_doc_uses_word_fallback_when_libreoffice_is_unavailable(self):
        docx_buffer = BytesIO()
        document_xml = '''
        <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:body>
            <w:p><w:r><w:t>课题研究内容：建立规则引擎并生成勘察报告。</w:t></w:r></w:p>
            <w:p><w:r><w:t>考核指标：完成一套报告生成工具。</w:t></w:r></w:p>
          </w:body>
        </w:document>
        '''
        with zipfile.ZipFile(docx_buffer, 'w') as archive:
            archive.writestr('word/document.xml', document_xml.encode('utf-8'))

        with patch('core.ai_analysis.os.path.exists', return_value=False), patch.object(
            AIAnalysisService,
            '_convert_doc_with_word',
            return_value=docx_buffer.getvalue(),
        ) as converter:
            extracted = AIAnalysisService._extract_doc_text(b'legacy-doc', '任务书.doc')

        converter.assert_called_once_with(b'legacy-doc', '任务书.doc')
        self.assertIn('建立规则引擎', extracted)
        self.assertIn('完成一套报告生成工具', extracted)

    @patch('core.ai_analysis.ai_service.analyze_document')
    def test_ai_reanalysis_syncs_metrics_without_losing_tracking_or_evidence(self, analyze_document):
        first_result = {
            'success': True,
            'result': '任务书考核指标',
            'structured_data': {'metrics': []},
            'metrics': [{
                'category': 'deliverable',
                'item_name': '老旧社区城市更新设计技术策略集成',
                'target_value': '1',
                'assessment_method': '征求意见稿',
                'planned_period': '',
                'deadline': '',
                'source_section': '考核指标',
                'source_page': '7',
                'notes': '',
                'sort_order': 1,
            }],
            'processing_time': 0.1,
            'api_used': 'TestAI',
            'confidence_score': 0.9,
        }
        analyze_document.return_value = first_result
        upload = SimpleUploadedFile('任务书.pdf', b'%PDF test', content_type='application/pdf')
        response = self.client.post(
            reverse('analyze_content', args=[self.project.project_id, 'output_metrics']),
            {'files': upload},
        )
        self.assertEqual(response.status_code, 302)
        metric = MetricsItem.objects.get()
        metric.current_value = '完成初稿'
        metric.progress_percent = 60
        metric.status = 'in_progress'
        metric.save()
        evidence = MetricEvidence.objects.create(
            metric=metric,
            relative_path='06_结题/策略集成稿.docx',
            display_name='策略集成稿.docx',
            created_by=self.admin,
        )

        second_result = dict(first_result)
        second_result['metrics'] = [dict(first_result['metrics'][0], target_value='2')]
        analyze_document.return_value = second_result
        upload = SimpleUploadedFile('任务书修订版.pdf', b'%PDF test 2', content_type='application/pdf')
        self.client.post(
            reverse('analyze_content', args=[self.project.project_id, 'output_metrics']),
            {'files': upload},
        )

        metric.refresh_from_db()
        self.assertEqual(MetricsItem.objects.count(), 1)
        self.assertEqual(metric.target_value, '2')
        self.assertEqual(metric.current_value, '完成初稿')
        self.assertEqual(metric.progress_percent, 60)
        self.assertEqual(metric.status, 'in_progress')
        self.assertTrue(MetricEvidence.objects.filter(pk=evidence.pk).exists())

    def test_metric_update_sets_completion_and_evidence_links_to_real_project_file(self):
        analysis = ProjectAnalysis.objects.create(
            project=self.project,
            analysis_type='output_metrics',
            file_name='任务书.pdf',
            analysis_result='指标',
        )
        metric = MetricsItem.objects.create(
            analysis=analysis,
            category='deliverable',
            item_name='代表性城市老旧社区更新专项方案研究',
            target_value='1',
        )
        with tempfile.TemporaryDirectory() as temp_root:
            project_folder = Path(temp_root) / 'project-folder'
            evidence_folder = project_folder / '06_结题'
            evidence_folder.mkdir(parents=True)
            (evidence_folder / '专项方案文本.pdf').write_bytes(b'%PDF evidence')
            self.project.directory_path = str(project_folder)
            self.project.save(update_fields=['directory_path'])

            update_response = self.client.post(
                reverse('update_metrics_item', args=[self.project.project_id, metric.id]),
                {
                    'status': 'completed',
                    'progress_percent': '95',
                    'current_value': '1',
                    'responsible_person': '张三',
                },
            )
            self.assertEqual(update_response.status_code, 302)
            metric.refresh_from_db()
            self.assertEqual(metric.status, 'completed')
            self.assertEqual(metric.progress_percent, 100)
            self.assertIsNotNone(metric.actual_completion_date)

            evidence_response = self.client.post(
                reverse('add_metric_evidence', args=[self.project.project_id, metric.id]),
                {'relative_path': '06_结题/专项方案文本.pdf', 'note': '专项方案成果'},
                HTTP_X_REQUESTED_WITH='XMLHttpRequest',
            )
            self.assertEqual(evidence_response.status_code, 200)
            linked = MetricEvidence.objects.get(metric=metric)
            self.assertEqual(linked.display_name, '专项方案文本.pdf')

    def test_manual_metric_uses_supplied_indicator_catalog_and_omits_schedule_fields(self):
        response = self.client.post(
            reverse('create_metrics_item', args=[self.project.project_id]),
            {
                'item_name': '软件著作权',
                'target_value': '2',
                'planned_period': '不应保存',
                'deadline': '2026-12-31',
                'responsible_person': '不应保存',
            },
        )

        self.assertEqual(response.status_code, 302)
        metric = MetricsItem.objects.get(item_name='软件著作权')
        self.assertEqual(metric.category, 'ip')
        self.assertEqual(metric.target_value, '2项')
        self.assertEqual(metric.assessment_method, '受理/授权')
        self.assertEqual(metric.planned_period, '')
        self.assertIsNone(metric.deadline)
        self.assertEqual(metric.responsible_person, '')

    def test_metric_maintenance_modal_has_scroll_layout_visible_return_and_catalog(self):
        analysis = ProjectAnalysis.objects.create(
            project=self.project,
            analysis_type='output_metrics',
            file_name='手动录入',
            analysis_result='指标',
        )
        MetricsItem.objects.create(
            analysis=analysis,
            category='ip',
            item_name='发明专利',
            target_value='1项',
            assessment_method='受理/授权',
        )

        with tempfile.TemporaryDirectory() as projects_root:
            with override_settings(PROJECTS_ROOT=Path(projects_root)):
                response = self.client.get(
                    reverse('project_detail', args=[self.project.project_id]),
                    {'tab': 'metrics-analysis'},
                )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '返回指标分析', count=3)
        self.assertContains(response, '核心期刊/SCI/EI论文')
        self.assertContains(response, '企业级/团体/地方/行业/国家标准')
        self.assertContains(response, '装备、工艺类')
        self.assertContains(response, '.modal-dialog-scrollable .modal-body { overflow-y: auto;')
        self.assertNotContains(response, 'name="planned_period"')
        self.assertNotContains(response, 'name="responsible_person"')
        self.assertNotContains(response, 'name="deadline"')

    def test_manual_metric_rejects_indicator_outside_catalog(self):
        response = self.client.post(
            reverse('create_metrics_item', args=[self.project.project_id]),
            {'item_name': '自定义临时指标', 'target_value': '1'},
        )

        self.assertEqual(response.status_code, 302)
        self.assertFalse(MetricsItem.objects.exists())

    def test_detail_renders_structured_research_and_metric_tracking_for_readonly_user(self):
        readonly = get_user_model().objects.create_user('analysis-readonly', password='StrongPass!234')
        research = ProjectAnalysis.objects.create(
            project=self.project,
            analysis_type='research_content',
            file_name='任务书.pdf',
            analysis_result='结构化研究内容',
            structured_data={
                'research_goal': '形成老旧社区城市更新技术策略',
                'research_sections': [{
                    'title': '老旧社区城市更新设计策略研究',
                    'description': '形成建筑修复集成技术研究成果',
                    'subitems': [],
                }],
                'methods': ['案例实践法'],
                'technical_route': [],
                'technical_difficulties': [],
                'innovations': ['形成老旧社区城市更新技术导则'],
                'milestones': [{'time_range': '2022.09-2022.12', 'stage_goal': '完成成果'}],
            },
        )
        analysis = ProjectAnalysis.objects.create(
            project=self.project,
            analysis_type='output_metrics',
            file_name='任务书.pdf',
            analysis_result='指标',
        )
        MetricsItem.objects.create(
            analysis=analysis,
            category='deliverable',
            item_name='城市更新视角下老旧社区改造的策略研究论文',
            target_value='1',
            assessment_method='发表',
            progress_percent=40,
        )
        self.assertEqual(research.project_id, self.project.pk)

        with tempfile.TemporaryDirectory() as projects_root:
            from .views import _get_project_folder_name

            expected_folder = Path(projects_root) / _get_project_folder_name(self.project)
            expected_folder.mkdir(parents=True)
            self.project.directory_path = str(expected_folder)
            self.project.save(update_fields=['directory_path'])
            self.client.force_login(readonly)
            with override_settings(PROJECTS_ROOT=Path(projects_root)):
                response = self.client.get(
                    reverse('project_detail', args=[self.project.project_id]),
                    {'tab': 'metrics-analysis'},
                )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '形成老旧社区城市更新技术策略')
        self.assertContains(response, '城市更新视角下老旧社区改造的策略研究论文')
        self.assertContains(response, '发表')
        self.assertContains(response, '40%')
        self.assertContains(response, '佐证材料')
        self.assertNotContains(response, '更新完成情况')


class ExpenseMonitorMonthlyLedgerTests(TestCase):
    COMPANY_A = TARGET_COMPANIES[0]
    COMPANY_B = TARGET_COMPANIES[1]

    def setUp(self):
        self.admin = get_user_model().objects.create_user(
            'expense-admin',
            password='StrongPass!234',
            is_staff=True,
        )
        self.client.force_login(self.admin)
        self.project = Project.objects.create(
            project_id='EXP-001',
            name='城市竖井关键技术研究',
            ownership='测试单位',
            managing_unit='测试归口单位',
            level='企业级',
            project_type='科研课题',
            role='牵头',
            start_year=2023,
            status='结题',
            project_lead='测试负责人',
            start_date=date(2023, 1, 1),
            planned_end_date=date(2024, 12, 31),
            actual_completion_date=date(2024, 12, 20),
            total_budget=Decimal('12.00'),
            directory_path='EXP-001',
        )
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.expense_path = Path(self.temp_dir.name) / '当前月支出.xlsx'
        self.settings_override = override_settings(EXPENSE_DATA_FILE=self.expense_path)
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)

    @staticmethod
    def _workbook_bytes(rows, valid=True):
        workbook = openpyxl.Workbook()
        intro = workbook.active
        intro.title = '说明'
        intro.append(['这是说明页'])
        detail = workbook.create_sheet('本月明细')
        if valid:
            detail.append([
                '公司名称',
                '总账科目',
                '利润中心名称',
                '科研课题文本描述',
                '本年累计借方金额',
                '期末余额',
            ])
            for row in rows:
                detail.append(row)
        else:
            detail.append(['公司名称', '科研课题文本描述', '期末余额'])
            detail.append(['错误公司', '错误课题', 100])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _write_workbook(self, rows):
        self.expense_path.write_bytes(self._workbook_bytes(rows))

    def _base_rows(self, company_b_amount=70000):
        return [
            [self.COMPANY_A, 6606, f'{self.COMPANY_A}-本部', '城市竖井关键技术研究-专项', 80000, 99999999],
            [self.COMPANY_B, '6606', f'{self.COMPANY_B}-本部', '城市竖井关键技术研究（地下空间公司自筹）', company_b_amount, 88888888],
            [self.COMPANY_A, 6401, f'{self.COMPANY_A}设计二院', '城市竖井关键技术研究', 999000, 77777777],
            [self.COMPANY_A, 660601, f'{self.COMPANY_A}-本部', '城市竖井关键技术研究', 555000, 77777777],
            [self.COMPANY_B, 6606, '天津地铁7号线项目', '尚未录入系统的课题（公司自筹）', 20000, 66666666],
        ]

    def test_analysis_filters_6606_merges_suffixes_and_compares_budget(self):
        self._write_workbook(self._base_rows())

        result = analyze_expense_workbook(self.expense_path, [self.project], threshold=0.85)

        self.assertEqual(result['sheet_name'], '本月明细')
        self.assertEqual(result['source_row_count'], 5)
        self.assertEqual(result['filtered_rows'], 3)
        self.assertEqual(result['ignored_account_rows'], 2)
        self.assertEqual(result['matched_project_total'], 1)
        self.assertEqual(len(result['unmatched_rows']), 1)
        self.assertEqual(result['total_expense_sum'], Decimal('17'))

        matched = result['project_rows'][0]
        self.assertEqual(matched['project_id'], self.project.project_id)
        self.assertEqual(matched['total'], Decimal('15'))
        self.assertEqual(matched['row_count'], 2)
        self.assertEqual(matched['variant_count'], 2)
        self.assertTrue(matched['over_budget'])
        self.assertEqual(matched['over_amount'], Decimal('3'))
        self.assertEqual(
            matched['profit_centers'],
            [f'{self.COMPANY_B}-本部', f'{SHORT_COMPANY_NAME}-本部'],
        )

        company_summary = {item['company']: item for item in result['company_summaries']}
        self.assertEqual(company_summary[SHORT_COMPANY_NAME]['total'], Decimal('8'))
        self.assertEqual(company_summary[self.COMPANY_B]['total'], Decimal('9'))
        self.assertEqual(company_summary[self.COMPANY_B]['matched_total'], Decimal('7'))
        self.assertEqual(company_summary[self.COMPANY_B]['unmatched_total'], Decimal('2'))

    def test_name_normalization_only_removes_funding_suffixes(self):
        base = '城市竖井关键技术研究'
        self.assertEqual(canonicalize_expense_description(f'{base}-专项'), base)
        self.assertEqual(canonicalize_expense_description(f'{base}（地下空间公司自筹）'), base)
        self.assertEqual(
            normalized_expense_description(f'{base}（地下空间公司自筹）'),
            normalized_expense_description(base),
        )
        self.assertEqual(
            canonicalize_expense_description('桩基技术研究（SPHC,竹节桩）'),
            '桩基技术研究(SPHC,竹节桩)',
        )

    def test_two_month_snapshots_flag_spending_after_completion(self):
        self._write_workbook(self._base_rows(company_b_amount=70000))
        first_response = self.client.get(reverse('expense_monitor'))

        self.assertEqual(first_response.status_code, 200)
        self.assertFalse(first_response.context['has_comparison'])
        self.assertEqual(ExpenseImport.objects.filter(format_version=EXPENSE_FORMAT_VERSION).count(), 1)
        self.assertEqual(ExpenseSnapshot.objects.count(), 1)

        self._write_workbook(self._base_rows(company_b_amount=90000))
        second_response = self.client.get(reverse('expense_monitor'))

        self.assertEqual(second_response.status_code, 200)
        self.assertTrue(second_response.context['has_comparison'])
        self.assertEqual(ExpenseImport.objects.filter(format_version=EXPENSE_FORMAT_VERSION).count(), 2)
        self.assertEqual(ExpenseSnapshot.objects.count(), 2)
        self.assertEqual(len(second_response.context['growth_alerts']), 1)
        alert = second_response.context['growth_alerts'][0]
        self.assertEqual(alert['project_id'], self.project.project_id)
        self.assertEqual(alert['previous_total'], Decimal('15'))
        self.assertEqual(alert['current_total'], Decimal('17'))
        self.assertEqual(alert['increase'], Decimal('2'))

    def test_invalid_upload_does_not_replace_current_workbook(self):
        self._write_workbook(self._base_rows())
        original_bytes = self.expense_path.read_bytes()
        invalid_upload = SimpleUploadedFile(
            '错误格式.xlsx',
            self._workbook_bytes([], valid=False),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        response = self.client.post(
            reverse('expense_import'),
            {'expense_file': invalid_upload},
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.expense_path.read_bytes(), original_bytes)

    def test_rendered_page_explains_new_account_and_amount_scope(self):
        self._write_workbook(self._base_rows())

        response = self.client.get(reverse('expense_monitor'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '总账科目6606')
        self.assertContains(response, '本年累计借方金额')
        self.assertContains(response, '利润中心名称')
        self.assertContains(response, SHORT_COMPANY_NAME)
        self.assertContains(response, self.COMPANY_B)
        self.assertNotContains(response, self.COMPANY_A)
        self.assertNotContains(response, '归口单位分组数')


class FundingCategoryTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('funding-admin', password='StrongPass!234', is_staff=True)
        self.client.force_login(self.user)
        common = dict(ownership='西勘院', managing_unit='测试单位', level='公司级', project_type='应用研究', role='牵头', start_year=2026, status='在研', start_date=date(2026, 1, 1))
        Project.objects.create(project_id='SPECIAL-1', name='专项课题', funding_category='special', **common)
        Project.objects.create(project_id='SELF-1', name='自筹课题', funding_category='self_funded', **common)

    def test_project_and_progress_pages_are_scoped(self):
        response = self.client.get(reverse('self_funded_project_list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '自筹课题')
        self.assertNotContains(response, '专项课题')
        progress = self.client.get(reverse('self_funded_progress_monitor'))
        self.assertEqual(progress.status_code, 200)
        self.assertContains(progress, '企业全自筹课题进度监控')

    def test_project_form_and_assistant_accept_category(self):
        form = ProjectForm(instance=Project.objects.get(project_id='SELF-1'))
        self.assertIn(('self_funded', '企业全自筹课题'), list(form.fields['funding_category'].choices))
        result = answer_project_question('查询所有课题', use_ai=False, funding_category='self_funded')
        self.assertTrue(result['ok'])
        self.assertEqual(result['total'], 1)
        self.assertEqual(result['rows'][0]['project'].project_id, 'SELF-1')


class ProjectIdWithSlashUrlTests(TestCase):
    """课题编号包含 /（如企业标准编号 QB/ZJXK0001-2021）时，路由仍可正常生成与访问。"""

    SLASH_ID = 'QB/ZJXK0001-2021'

    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user('slash-admin', password='StrongPass!234', is_staff=True)
        self.client.force_login(self.admin)
        self.project = Project.objects.create(
            project_id=self.SLASH_ID,
            name='带斜杠编号的课题',
            ownership='西勘院',
            managing_unit='测试单位',
            level='公司级',
            project_type='应用研究',
            role='牵头',
            start_year=2021,
            status='在研',
            directory_path='',
        )

    def test_detail_url_reverses_with_slash(self):
        self.assertEqual(
            reverse('project_detail', args=[self.SLASH_ID]),
            f'/project/{self.SLASH_ID}/',
        )

    def test_list_page_renders_slash_id_project(self):
        response = self.client.get(reverse('project_list'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'/project/{self.SLASH_ID}/')

    def test_detail_and_file_tree_pages_resolve(self):
        from .views import create_project_directory_structure

        with tempfile.TemporaryDirectory() as temp_dir:
            with override_settings(PROJECTS_ROOT=Path(temp_dir)):
                create_project_directory_structure(self.project)
                detail_response = self.client.get(reverse('project_detail', args=[self.SLASH_ID]))
                file_tree_response = self.client.get(reverse('get_file_tree', args=[self.SLASH_ID]))

        self.assertEqual(detail_response.status_code, 200)
        self.assertEqual(file_tree_response.status_code, 200)
